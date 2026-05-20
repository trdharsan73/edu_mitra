# ── SQLite Compatibility Override for Streamlit Cloud ──────────────────────────
try:
    __import__('pysqlite3')
    import sys
    sys.modules['sqlite3'] = sys.modules.pop('pysqlite3')
except ImportError:
    pass

import streamlit as st
import os
import tempfile
import json
import io
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_community.embeddings import HuggingFaceEmbeddings
import google.generativeai as genai
from pptx import Presentation

# ── Bootstrap ─────────────────────────────────────────────────────────────────
load_dotenv(override=True)

# Helper to fetch and configure key
def get_api_key():
    # 1. Try Streamlit Secrets
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    # 2. Try OS Environment
    return os.getenv("GEMINI_API_KEY")

api_key = get_api_key()
if api_key:
    genai.configure(api_key=api_key)

st.set_page_config(
    page_title="EduMitra – AI Tutor",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Theme toggle (must be before CSS injection) ───────────────────────────────
if "dark_mode" not in st.session_state:
    st.session_state.dark_mode = False
dark = st.session_state.dark_mode

# ── Design System CSS (theme-aware) ─────────────────────────────────────────
_t = {
    # ── Dark backgrounds: warm charcoal layered system ──
    "app_bg":        "#0b0c10" if dark else "#f9f9f9",
    "sidebar_bg":    "#111319" if dark else "#f3f3f3",
    "sidebar_border":"#232836" if dark else "#bccac0",
    "card_bg":       "#161922" if dark else "#ffffff",
    "card_border":   "#232836" if dark else "#e2e2e2",
    "right_bg":      "#111319" if dark else "#f3f3f3",
    # ── Typography ──
    "text_primary":  "#f3f4f6" if dark else "#1a1c1c",
    "text_secondary":"#9ca3af" if dark else "#3d4a42",
    "text_label":    "#9ca3af" if dark else "#3d4a3d",
    # ── Brand green: emerald system ──
    "primary":       "#68dba9" if dark else "#006948",
    "primary_hover": "#85f8c4" if dark else "#005137",
    "primary_light": "rgba(104,219,169,0.1)" if dark else "rgba(0,105,72,0.08)",
    # ── Chat bubbles ──
    "ai_bubble_bg":  "#161922" if dark else "#ffffff",
    "ai_bubble_bdr": "#232836" if dark else "#e2e2e2",
    "user_bubble_bg":"#1e293b" if dark else "#e2e2e2",
    "user_bubble_bdr":"transparent" if dark else "transparent",
    # ── Inputs / form fields ──
    "input_bg":      "#161922" if dark else "#ffffff",
    "input_border":  "#232836" if dark else "#e2e2e2",
    # ── Misc ──
    "hr_color":      "#232836" if dark else "#e2e2e2",
    "scrollbar":     "#232836" if dark else "#dadada",
    "badge_bg":      "rgba(104,219,169,0.1)" if dark else "rgba(0,105,72,0.05)",
    "badge_border":  "rgba(104,219,169,0.2)" if dark else "rgba(0,105,72,0.15)",
    "select_bg":     "#161922" if dark else "#ffffff",
}

st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@500;600;700;800&family=Inter:wght@400;500;600;700&display=swap');
@import url('https://fonts.googleapis.com/css2?family=Material+Symbols+Outlined:wght,FILL@100..700,0..1&display=swap');

*, *::before, *::after {{ box-sizing: border-box; }}
html, body, .stApp {{
    font-family: 'Inter', sans-serif;
    background-color: {_t['app_bg']} !important;
    color: {_t['text_primary']};
    transition: background-color 0.3s ease, color 0.3s ease;
}}

h1, h2, h3, h4, h5, h6 {{
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-weight: 700 !important;
    letter-spacing: -0.02em !important;
}}

#MainMenu, footer, header {{ visibility: hidden; }}
.stDeployButton {{ display: none; }}
.block-container {{ padding: 0 !important; max-width: 100% !important; }}

/* Sidebar */
section[data-testid="stSidebar"] {{
    background-color: {_t['sidebar_bg']} !important;
    border-right: 1px solid {_t['sidebar_border']} !important;
    padding: 0 !important;
    transition: background-color 0.3s ease;
}}
section[data-testid="stSidebar"] > div:first-child {{ padding: 0 !important; }}

/* Selectbox & Labels */
.stSelectbox label, .stCheckbox label, .stFileUploader label {{
    color: {_t['text_label']} !important;
    font-size: 13px !important;
    font-weight: 600 !important;
    letter-spacing: 0.05em !important;
    text-transform: uppercase !important;
}}
.stSelectbox > div > div {{
    background: {_t['select_bg']} !important;
    border: 1px solid {_t['card_border']} !important;
    border-radius: 12px !important;
    color: {_t['text_primary']} !important;
    font-size: 14px !important;
    padding: 2px 4px !important;
    transition: border-color 0.2s ease;
}}
.stSelectbox > div > div:focus-within {{
    border-color: {_t['primary']} !important;
}}

/* Checkbox Toggle Switch Styling */
.stCheckbox span {{ color: {_t['text_primary']} !important; font-size: 14px !important; }}

/* Buttons */
.stButton > button {{
    background-color: {_t['primary']} !important;
    color: #ffffff !important;
    border: none !important;
    border-radius: 9999px !important;
    font-size: 14px !important;
    font-weight: 600 !important;
    padding: 10px 24px !important;
    width: 100% !important;
    transition: all 0.2s ease-in-out !important;
    box-shadow: 0 4px 12px rgba(0, 105, 72, 0.15) !important;
}}
.stButton > button:hover {{
    background-color: {_t['primary_hover']} !important;
    transform: translateY(-1px) !important;
    box-shadow: 0 6px 16px rgba(0, 105, 72, 0.25) !important;
}}
.stButton > button:active {{ transform: scale(0.98) !important; }}

/* Chat Messages */
div[data-testid="stChatMessage"] {{
    background-color: transparent !important;
    border: none !important;
    padding: 6px 0 !important;
}}
div[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-assistant"]) > div:last-child {{
    background: {_t['ai_bubble_bg']} !important;
    border: 1px solid {_t['ai_bubble_bdr']} !important;
    border-radius: 20px 20px 20px 4px !important;
    padding: 18px 24px !important;
    box-shadow: 0 4px 20px -8px rgba(0,0,0,0.06), 0 2px 6px -4px rgba(0,0,0,0.03) !important;
}}
div[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) > div:last-child {{
    background: {_t['user_bubble_bg']} !important;
    border: none !important;
    border-radius: 20px 20px 4px 20px !important;
    padding: 14px 20px !important;
    box-shadow: 0 2px 8px -2px rgba(0,0,0,0.04) !important;
}}
div[data-testid="stChatMessageContent"] p {{
    color: {_t['text_primary']} !important;
    font-size: 15px !important;
    line-height: 1.65 !important;
}}

/* Customizing Streamlit's Chat Message Avatars */
div[data-testid="stChatMessage"] div[data-testid="chatAvatarIcon-assistant"] {{
    background-color: {_t['badge_bg']} !important;
    color: {_t['primary']} !important;
    border-radius: 50% !important;
}}
div[data-testid="stChatMessage"] div[data-testid="chatAvatarIcon-user"] {{
    background-color: {_t['user_bubble_bg']} !important;
    color: {_t['text_primary']} !important;
    border-radius: 50% !important;
}}

/* Chat Input */
div[data-testid="stChatInput"] {{
    background: {_t['input_bg']} !important;
    border: 1px solid {_t['input_border']} !important;
    border-radius: 9999px !important;
    box-shadow: 0 8px 32px -8px rgba(0,0,0,0.08) !important;
    padding: 6px 12px !important;
    transition: focus-within 0.25s ease;
}}
div[data-testid="stChatInput"]:focus-within {{
    border-color: {_t['primary']} !important;
    box-shadow: 0 8px 32px -8px rgba(0, 105, 72, 0.12) !important;
}}
div[data-testid="stChatInput"] textarea {{
    color: {_t['text_primary']} !important;
    font-size: 15px !important;
    background: transparent !important;
}}
div[data-testid="stChatInput"] button {{
    background-color: {_t['primary']} !important;
    border-radius: 50% !important;
}}

/* File Uploader */
div[data-testid="stFileUploader"] > div {{
    background: {_t['card_bg']} !important;
    border: 1.5px dashed {_t['card_border']} !important;
    border-radius: 14px !important;
    padding: 20px !important;
    transition: border-color 0.2s ease;
}}
div[data-testid="stFileUploader"] > div:hover {{
    border-color: {_t['primary']} !important;
}}
div[data-testid="stFileUploader"] p, div[data-testid="stFileUploader"] span {{
    color: {_t['text_secondary']} !important;
}}

/* Alert */
div[data-testid="stAlert"] {{
    background: {_t['badge_bg']} !important;
    border: 1px solid {_t['badge_border']} !important;
    border-radius: 14px !important;
    color: {_t['text_primary']} !important;
}}

/* Divider */
hr {{ border-color: {_t['hr_color']} !important; }}

/* Spinner */
div[data-testid="stSpinner"] p {{ color: {_t['primary']} !important; }}

/* Scrollbar */
::-webkit-scrollbar {{ width: 6px; }}
::-webkit-scrollbar-track {{ background: transparent; }}
::-webkit-scrollbar-thumb {{ background: {_t['scrollbar']}; border-radius: 9999px; }}
</style>
""", unsafe_allow_html=True)


# ── Session State ──────────────────────────────────────────────────────────────
if "gemini_files" not in st.session_state:
    st.session_state.gemini_files = {}
if "messages" not in st.session_state:
    st.session_state.messages = []
if "generate_ppt" not in st.session_state:
    st.session_state.generate_ppt = False


# ── RAG Vector Store ───────────────────────────────────────────────────────────
@st.cache_resource
def load_vector_store():
    db_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "chroma_db")
    if os.path.exists(db_path):
        embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
        return Chroma(persist_directory=db_path, embedding_function=embeddings)
    return None

db = load_vector_store()


# ── AI Response ────────────────────────────────────────────────────────────────
def get_ai_response(user_prompt, class_level, subject, language, step_by_step, db):
    global api_key
    if not api_key:
        api_key = get_api_key()
        if api_key:
            genai.configure(api_key=api_key)
            
    if not api_key:
        return "⚠️ Please configure `GEMINI_API_KEY` in your Streamlit Secrets or `.env` file."

    db_context = ""
    if db:
        try:
            docs = db.similarity_search(user_prompt, k=4)
            db_context = "\n".join([d.page_content for d in docs])
        except Exception:
            pass

    step_txt = "Always give a clear, numbered step-by-step breakdown." if step_by_step else ""
    system = f"""You are EduMitra, an expert AI tutor for Indian school and college students.
Help a Class {class_level} student with {subject}. Respond in: {language}.
Use the DATABASE CONTEXT and any UPLOADED FILES to give grounded answers.
Format answers with headers, bullet points, and helpful emojis. {step_txt}

--- DATABASE CONTEXT ---
{db_context}
--- END CONTEXT ---"""

    try:
        model = genai.GenerativeModel("gemini-2.5-flash")
        parts = [system]
        for gf in st.session_state.gemini_files.values():
            parts.append(gf)
        history = "\n".join([
            f"{m['role'].capitalize()}: {m['content']}"
            for m in st.session_state.messages[-6:]
        ])
        if history:
            parts.append(f"Recent conversation:\n{history}")
        parts.append(f"Student: {user_prompt}")
        return model.generate_content(parts).text
    except Exception as e:
        return f"⚠️ Error: {e}"


# ── PPT Helper ─────────────────────────────────────────────────────────────────
def create_pptx(slides_data):
    prs = Presentation()
    for s in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for b in s.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.level = 0
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# SIDEBAR
# ══════════════════════════════════════════════════════════════════════════════
with st.sidebar:
    # Brand + Theme Toggle
    toggle_icon = "☀️" if dark else "🌙"
    toggle_label = "Light Mode" if dark else "Dark Mode"
    brand_border = _t['sidebar_border']
    brand_color  = _t['primary']
    sub_color    = _t['text_secondary']
    st.markdown(f"""
    <div style="padding:24px 24px 16px; border-bottom:1px solid {brand_border};
                display:flex; flex-direction:column; gap:4px; margin-bottom:12px;">
        <h1 style="margin:0; font-size:26px; font-weight:800; color:{brand_color}; letter-spacing:-0.03em;">EduMitra</h1>
        <p style="margin:0; font-size:11px; font-weight:600; color:{sub_color}; text-transform:uppercase; letter-spacing:0.1em;">Intelligent Growth</p>
    </div>
    """, unsafe_allow_html=True)

    # Dark/Light toggle button
    if st.button(f"{toggle_icon}  {toggle_label}", key="theme_toggle"):
        st.session_state.dark_mode = not st.session_state.dark_mode
        st.rerun()

    # Nav links (visual only)
    nav_items = [
        ("school", "AI Tutor"),
        ("insights", "Performance Dashboard"),
    ]
    st.markdown(f'<div style="padding:8px 16px 8px;">', unsafe_allow_html=True)
    for icon, label in nav_items:
        active = label == "AI Tutor"
        bg     = _t['primary_light'] if active else "transparent"
        color  = _t['primary'] if active else _t['text_secondary']
        weight = "600" if active else "400"
        fill   = "1" if active else "0"
        st.markdown(f"""
        <div style="display:flex;align-items:center;gap:12px;padding:10px 14px;
                    border-radius:12px;background:{bg};margin-bottom:2px;cursor:pointer;
                    transition:background 0.2s ease, transform 0.2s ease;">
            <span class="material-symbols-outlined"
                  style="font-size:20px;color:{color};font-variation-settings:'FILL' {fill};">{icon}</span>
            <span style="font-size:14px;font-weight:{weight};color:{color};">{label}</span>
        </div>
        """, unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown(f'<div style="padding:0 16px;">', unsafe_allow_html=True)
    st.markdown("---")

    # Tutor Preferences
    st.markdown(f'<p style="font-size:11px;font-weight:700;color:{_t["text_label"]};text-transform:uppercase;letter-spacing:0.08em;margin-bottom:12px;">Tutor Settings</p>', unsafe_allow_html=True)

    class_level = st.selectbox("Class Level", ["8", "9", "10", "11", "12", "College"])
    subject     = st.selectbox("Subject", ["Mathematics", "Physics", "Chemistry", "Biology", "Science", "General"])
    language    = st.selectbox("Language", ["English", "Tamil", "English + Tamil (Tamlish)"])
    step_by_step = st.checkbox("Step-by-step explanations", value=True)

    st.markdown("---")
    st.markdown(f'<p style="font-size:11px;color:{_t["text_secondary"]};line-height:1.5;">Powered by Gemini 2.5 Flash · 48,000+ NCERT chunks</p>', unsafe_allow_html=True)

    # Footer links
    st.markdown(f"""
    <div style="display:flex;justify-content:space-between;padding-top:12px;border-top:1px solid {_t['card_border']};margin-top:8px;">
        <a href="#" style="color:{_t['text_secondary']};text-decoration:none;font-size:12px;font-weight:500;
                           display:inline-flex;align-items:center;gap:4px;transition:color 0.2s ease;">
            <span class="material-symbols-outlined" style="font-size:16px;">settings</span> Settings
        </a>
        <a href="#" style="color:{_t['text_secondary']};text-decoration:none;font-size:12px;font-weight:500;
                           display:inline-flex;align-items:center;gap:4px;transition:color 0.2s ease;">
            <span class="material-symbols-outlined" style="font-size:16px;">help</span> Help
        </a>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("</div>", unsafe_allow_html=True)

    # Upgrade CTA
    st.markdown(f"""
    <div style="background:{_t['card_bg']}; border: 1px solid {_t['card_border']};
                border-radius: 12px; padding: 16px; margin: 12px 16px 8px;
                box-shadow: 0 4px 12px rgba(0,0,0,0.03); display: flex; flex-direction: column; gap: 10px;">
        <div style="display: flex; align-items: center; gap: 8px; color: {_t['primary']};">
            <span class="material-symbols-outlined" style="font-size: 20px; font-variation-settings: 'FILL' 1;">workspace_premium</span>
            <h3 style="margin: 0; font-size: 14px; font-weight: 700; color: {_t['text_primary']};">EduMitra Pro</h3>
        </div>
        <p style="margin: 0; font-size: 12px; color: {_t['text_secondary']}; line-height: 1.5;">
            Unlock infinite PDF uploads and priority AI processing.
        </p>
    </div>
    """, unsafe_allow_html=True)

    st.markdown('<div style="padding: 0 16px 20px;">', unsafe_allow_html=True)
    st.button("✨ Go Premium", key="go_premium")
    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# MAIN LAYOUT  (Chat | Right Panel)
# ══════════════════════════════════════════════════════════════════════════════
chat_col, right_col = st.columns([7, 3], gap="large")


# ── RIGHT PANEL ───────────────────────────────────────────────────────────────
with right_col:
    st.markdown(f"""
    <div style="padding: 24px 0 0; display: flex; align-items: center; gap: 8px; margin-bottom: 16px; border-bottom: 1px solid {_t['card_border']}; padding-bottom: 16px;">
        <span class="material-symbols-outlined" style="color: {_t['primary']}; font-size: 22px;">library_books</span>
        <h3 style="margin: 0; font-size: 18px; font-weight: 700; color: {_t['text_primary']};">Knowledge Base</h3>
    </div>
    """, unsafe_allow_html=True)

    # File uploader
    uploaded_files = st.file_uploader(
        "Upload PDFs or images",
        type=["pdf", "png", "jpg", "jpeg"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )

    if st.button("⬆️  Upload & Process"):
        if not api_key:
            api_key = get_api_key()
            if api_key:
                genai.configure(api_key=api_key)
                
        if not api_key:
            st.error("Add GEMINI_API_KEY to your Streamlit Secrets or .env file.")
        elif not uploaded_files:
            st.warning("Select at least one file first.")
        else:
            with st.spinner("Processing files…"):
                for uf in uploaded_files:
                    if uf.name not in st.session_state.gemini_files:
                        ext = "." + uf.name.split(".")[-1]
                        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                            tmp.write(uf.getvalue())
                            tmp_path = tmp.name
                        try:
                            gf = genai.upload_file(tmp_path, display_name=uf.name)
                            st.session_state.gemini_files[uf.name] = gf
                            st.success(f"✅ {uf.name}")
                        except Exception as e:
                            st.error(f"Failed: {e}")
                        finally:
                            os.remove(tmp_path)

    # Active sources list
    if st.session_state.gemini_files:
        st.markdown(f'<p style="font-size:11px;font-weight:700;color:{_t["text_label"]};text-transform:uppercase;letter-spacing:0.08em;margin:20px 0 10px;">Active Sources</p>', unsafe_allow_html=True)
        for name in st.session_state.gemini_files:
            ext = name.split(".")[-1].upper()
            icon_color = "#E5252A" if ext == "PDF" else _t['primary']
            mat_icon = "picture_as_pdf" if ext == "PDF" else "image"
            st.markdown(f"""
            <div style="display:flex;align-items:center;gap:10px;padding:10px 12px;
                        background:{_t['card_bg']};border:1px solid {_t['card_border']};
                        border-radius:10px;margin-bottom:6px;
                        box-shadow:0 2px 6px rgba(0,0,0,0.03);
                        transition:transform 0.15s ease;">
                <span class="material-symbols-outlined" style="color:{icon_color};font-size:22px;">{mat_icon}</span>
                <div style="min-width:0;flex:1;">
                    <p style="font-size:13px;font-weight:600;color:{_t['text_primary']};margin:0;
                               white-space:nowrap;overflow:hidden;text-overflow:ellipsis;">{name}</p>
                    <p style="font-size:11px;color:{_t['text_secondary']};margin:0;">Loaded · {ext}</p>
                </div>
            </div>
            """, unsafe_allow_html=True)

        if st.button("🗑️  Clear All Sources"):
            for gf in st.session_state.gemini_files.values():
                try:
                    genai.delete_file(gf.name)
                except Exception:
                    pass
            st.session_state.gemini_files = {}
            st.rerun()

    st.markdown("---")

    # PPT Studio
    st.markdown(f"""
    <div style="margin-top: 24px; background: linear-gradient(135deg, {_t['card_bg']}, {_t['right_bg']}); border-radius: 12px; padding: 20px; border: 1px solid {_t['card_border']}; box-shadow: 0 4px 12px rgba(0,0,0,0.03); position: relative; overflow: hidden; margin-bottom: 16px;">
        <div style="display: flex; align-items: center; gap: 8px; color: {_t['primary']}; margin-bottom: 8px;">
            <span class="material-symbols-outlined" style="font-size: 20px; font-variation-settings: 'FILL' 1;">co_present</span>
            <h3 style="margin: 0; font-size: 15px; font-weight: 700; color: {_t['text_primary']};">Presentation Studio</h3>
        </div>
        <p style="margin: 0; font-size: 13px; color: {_t['text_secondary']}; line-height: 1.5;">
            Convert current chat context and documents into a structured slide deck.
        </p>
    </div>
    """, unsafe_allow_html=True)

    if st.button("✨  Create Summary PPT"):
        st.session_state.generate_ppt = True

    if st.session_state.generate_ppt:
        with st.spinner("Crafting slides…"):
            try:
                model = genai.GenerativeModel("gemini-2.5-flash")
                prompt = """Summarize key concepts into a presentation.
Output ONLY valid JSON (no markdown fences):
[{"title": "...", "bullets": ["...", "..."]}]
Create 5–10 educational slides."""
                parts = [prompt]
                for gf in st.session_state.gemini_files.values():
                    parts.append(gf)
                hist = "\n".join([f"{m['role'].capitalize()}: {m['content']}" for m in st.session_state.messages])
                if hist:
                    parts.append(f"Conversation:\n{hist}")
                raw = model.generate_content(parts).text.strip()
                for marker in ["```json", "```"]:
                    raw = raw.replace(marker, "")
                slides = json.loads(raw)
                ppt_buf = create_pptx(slides)
                st.success("Presentation ready!")
                st.download_button(
                    "📥  Download .pptx",
                    data=ppt_buf,
                    file_name="EduMitra_Summary.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            except Exception as e:
                st.error(f"Failed: {e}")
        st.session_state.generate_ppt = False


# ── CHAT PANEL ────────────────────────────────────────────────────────────────
with chat_col:
    # Top bar
    ncert_loaded = bool(db)
    sources_loaded = bool(st.session_state.gemini_files)
    badge_bg_color = _t['badge_bg']
    badge_border_color = _t['badge_border']
    badge_text_color = _t['primary'] if (ncert_loaded or sources_loaded) else _t['text_secondary']
    badge_text  = "NCERT Context Loaded" if ncert_loaded else "No RAG Context"

    st.markdown(f"""
    <div style="display:flex;justify-content:space-between;align-items:center;
                padding:20px 4px 16px;border-bottom:1px solid {_t['card_border']};margin-bottom:20px;">
        <h2 style="margin:0;font-size:22px;font-weight:700;color:{_t['primary']};">AI Tutor</h2>
        <span style="display:inline-flex;align-items:center;gap:6px;background:{badge_bg_color};
                     border:1px solid {badge_border_color};border-radius:20px;padding:5px 12px;
                     font-size:12px;font-weight:600;color:{badge_text_color};">
            ● {badge_text}
        </span>
    </div>
    """, unsafe_allow_html=True)

    # Welcome message
    if not st.session_state.messages:
        st.markdown(f"""
        <div style="background:{_t['card_bg']}; border:1px solid {_t['card_border']};
                    border-radius:16px; padding:28px; margin-bottom:24px;
                    box-shadow:0 4px 24px -8px rgba(0,0,0,0.06); transition:transform 0.3s ease;">
            <div style="display:flex; align-items:flex-start; gap:16px; margin-bottom:20px;">
                <div style="width:48px; height:48px; border-radius:50%; background:{_t['primary_light']};
                            display:flex; align-items:center; justify-content:center; flex-shrink:0;">
                    <span class="material-symbols-outlined" style="color:{_t['primary']}; font-size:28px;">smart_toy</span>
                </div>
                <div>
                    <h2 style="margin:0 0 6px; font-size:24px; font-weight:700; color:{_t['text_primary']};">Hello! I'm EduMitra.</h2>
                    <p style="margin:0; font-size:15px; color:{_t['text_secondary']}; line-height:1.6;">
                        Your personalized AI tutor. I notice we are focusing on <strong style="color:{_t['primary']}; font-weight:600;">Class {class_level} {subject}</strong> today. How can I assist you with your studies?
                    </p>
                </div>
            </div>
            <div style="display:flex; flex-wrap:wrap; gap:12px; margin-top:20px; border-top:1px solid {_t['card_border']}; padding-top:20px;">
                <span style="background:{_t['app_bg']}; border:1px solid {_t['card_border']}; border-radius:9999px;
                             padding:8px 16px; font-size:13px; color:{_t['text_primary']}; font-weight:500;
                             display:inline-flex; align-items:center; gap:8px;">
                    <span class="material-symbols-outlined" style="color:{_t['primary']}; font-size:18px;">step</span>
                    Explain step-by-step
                </span>
                <span style="background:{_t['app_bg']}; border:1px solid {_t['card_border']}; border-radius:9999px;
                             padding:8px 16px; font-size:13px; color:{_t['text_primary']}; font-weight:500;
                             display:inline-flex; align-items:center; gap:8px;">
                    <span class="material-symbols-outlined" style="color:{_t['primary']}; font-size:18px;">quiz</span>
                    Generate practice questions
                </span>
                <span style="background:{_t['app_bg']}; border:1px solid {_t['card_border']}; border-radius:9999px;
                             padding:8px 16px; font-size:13px; color:{_t['text_primary']}; font-weight:500;
                             display:inline-flex; align-items:center; gap:8px;">
                    <span class="material-symbols-outlined" style="color:{_t['primary']}; font-size:18px;">translate</span>
                    Answer in Tanglish
                </span>
            </div>
        </div>
        """, unsafe_allow_html=True)

    # Chat history
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    # Chat input
    if user_input := st.chat_input("Ask EduMitra anything…"):
        st.session_state.messages.append({"role": "user", "content": user_input})
        with st.chat_message("user"):
            st.markdown(user_input)

        with st.chat_message("assistant"):
            with st.spinner("EduMitra is thinking…"):
                reply = get_ai_response(user_input, class_level, subject, language, step_by_step, db)
            st.markdown(reply)

        st.session_state.messages.append({"role": "assistant", "content": reply})
