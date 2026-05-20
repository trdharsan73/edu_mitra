# ── SQLite Compatibility Override for Streamlit Cloud ──────────────────────────
try:
    __import__('pysqlite3')
    import sys
    sys.modules['sqlite3'] = sys.modules.pop('pysqlite3')
except ImportError:
    pass

import streamlit as st
import os
import tempfile
import json
import io
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_community.embeddings import HuggingFaceEmbeddings
import google.generativeai as genai
from pptx import Presentation

# ── Bootstrap ─────────────────────────────────────────────────────────────────
load_dotenv(override=True)
api_key = os.getenv("GEMINI_API_KEY")
if api_key:
    genai.configure(api_key=api_key)

st.set_page_config(
    page_title="EduMitra – AI Tutor",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Theme toggle (must be before CSS injection) ───────────────────────────────
if "dark_mode" not in st.session_state:
    st.session_state.dark_mode = False
dark = st.session_state.dark_mode

# ── Design System CSS (theme-aware) ─────────────────────────────────────────
_t = {
    # ── Dark backgrounds: warm charcoal layered system ──
    "app_bg":        "#13151f" if dark else "#f8f9ff",
    "sidebar_bg":    "#1a1d2e" if dark else "#f8f9ff",
    "sidebar_border":"#2a2f45" if dark else "#dce9ff",
    "card_bg":       "#1e2235" if dark else "#ffffff",
    "card_border":   "#2e3450" if dark else "#dce9ff",
    "right_bg":      "#1a1d2e" if dark else "#eff4ff",
    # ── Typography: high-contrast slate scale ──
    "text_primary":  "#e8eaf4" if dark else "#0b1c30",
    "text_secondary":"#7c85a2" if dark else "#6d7b6c",
    "text_label":    "#9ba3bf" if dark else "#3d4a3d",
    # ── Brand green: vivid emerald for dark mode ──
    "primary":       "#4ade80" if dark else "#006e2f",
    "primary_hover": "#22c55e" if dark else "#005321",
    "primary_light": "rgba(74,222,128,0.12)" if dark else "rgba(0,110,47,0.07)",
    # ── Chat bubbles: distinctly tinted layers ──
    "ai_bubble_bg":  "#18261e" if dark else "#f0fdf4",
    "ai_bubble_bdr": "rgba(74,222,128,0.22)" if dark else "#bbf7d0",
    "user_bubble_bg":"#1c2240" if dark else "#e5eeff",
    "user_bubble_bdr":"rgba(120,140,255,0.25)" if dark else "#dce9ff",
    # ── Inputs / form fields ──
    "input_bg":      "#1e2235" if dark else "#ffffff",
    "input_border":  "#2e3450" if dark else "#dce9ff",
    # ── Misc ──
    "hr_color":      "#2a2f45" if dark else "#dce9ff",
    "scrollbar":     "#2a2f45" if dark else "#dce9ff",
    "badge_bg":      "#18261e" if dark else "#f0fdf4",
    "badge_border":  "rgba(74,222,128,0.3)" if dark else "#bbf7d0",
    "select_bg":     "#1e2235" if dark else "#ffffff",
}

st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

*, *::before, *::after {{ box-sizing: border-box; }}
html, body, .stApp {{
    font-family: 'Inter', sans-serif;
    background-color: {_t['app_bg']} !important;
    color: {_t['text_primary']};
    transition: background-color 0.3s ease, color 0.3s ease;
}}

#MainMenu, footer, header {{ visibility: hidden; }}
.stDeployButton {{ display: none; }}
.block-container {{ padding: 0 !important; max-width: 100% !important; }}

/* Sidebar */
section[data-testid="stSidebar"] {{
    background-color: {_t['sidebar_bg']} !important;
    border-right: 1px solid {_t['sidebar_border']} !important;
    padding: 0 !important;
    transition: background-color 0.3s ease;
}}
section[data-testid="stSidebar"] > div:first-child {{ padding: 0 !important; }}

/* Selectbox & Labels */
.stSelectbox label, .stCheckbox label, .stFileUploader label {{
    color: {_t['text_label']} !important;
    font-size: 13px !important;
    font-weight: 600 !important;
}}
.stSelectbox > div > div {{
    background: {_t['select_bg']} !important;
    border: 1px solid {_t['card_border']} !important;
    border-radius: 10px !important;
    color: {_t['text_primary']} !important;
    font-size: 14px !important;
}}

/* Checkbox */
.stCheckbox span {{ color: {_t['text_primary']} !important; }}

/* Buttons */
.stButton > button {{
    background-color: {_t['primary']} !important;
    color: #ffffff !important;
    border: none !important;
    border-radius: 10px !important;
    font-size: 14px !important;
    font-weight: 600 !important;
    padding: 10px 20px !important;
    width: 100% !important;
    transition: background 0.2s, transform 0.15s !important;
    box-shadow: 0 2px 8px rgba(0,110,47,0.2) !important;
}}
.stButton > button:hover {{
    background-color: {_t['primary_hover']} !important;
    transform: translateY(-1px) !important;
}}
.stButton > button:active {{ transform: scale(0.98) !important; }}

/* Chat Messages */
div[data-testid="stChatMessage"] {{
    background-color: transparent !important;
    border: none !important;
    padding: 4px 0 !important;
}}
div[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-assistant"]) > div:last-child {{
    background: {_t['ai_bubble_bg']} !important;
    border: 1px solid {_t['ai_bubble_bdr']} !important;
    border-radius: 20px 20px 20px 4px !important;
    padding: 16px 20px !important;
    box-shadow: 0 1px 4px rgba(0,0,0,0.08) !important;
}}
div[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) > div:last-child {{
    background: {_t['user_bubble_bg']} !important;
    border: 1px solid {_t['user_bubble_bdr']} !important;
    border-radius: 20px 20px 4px 20px !important;
    padding: 16px 20px !important;
}}
div[data-testid="stChatMessageContent"] p {{
    color: {_t['text_primary']} !important;
    font-size: 15px !important;
    line-height: 1.65 !important;
}}

/* Chat Input */
div[data-testid="stChatInput"] {{
    background: {_t['input_bg']} !important;
    border: 1.5px solid {_t['input_border']} !important;
    border-radius: 24px !important;
    box-shadow: 0 4px 20px rgba(0,0,0,0.08) !important;
    padding: 4px 8px !important;
}}
div[data-testid="stChatInput"] textarea {{
    color: {_t['text_primary']} !important;
    font-size: 15px !important;
    background: transparent !important;
}}
div[data-testid="stChatInput"] button {{
    background-color: {_t['primary']} !important;
    border-radius: 50% !important;
}}

/* File Uploader */
div[data-testid="stFileUploader"] > div {{
    background: {_t['card_bg']} !important;
    border: 1.5px dashed {_t['card_border']} !important;
    border-radius: 14px !important;
    padding: 16px !important;
}}
div[data-testid="stFileUploader"] p, div[data-testid="stFileUploader"] span {{
    color: {_t['text_secondary']} !important;
}}

/* Alert */
div[data-testid="stAlert"] {{
    background: {_t['badge_bg']} !important;
    border: 1px solid {_t['badge_border']} !important;
    border-radius: 12px !important;
    color: {_t['text_primary']} !important;
}}

/* Divider */
hr {{ border-color: {_t['hr_color']} !important; }}

/* Spinner */
div[data-testid="stSpinner"] p {{ color: {_t['primary']} !important; }}

/* Scrollbar */
::-webkit-scrollbar {{ width: 5px; }}
::-webkit-scrollbar-track {{ background: transparent; }}
::-webkit-scrollbar-thumb {{ background: {_t['scrollbar']}; border-radius: 10px; }}
</style>
""", unsafe_allow_html=True)


# ── Session State ──────────────────────────────────────────────────────────────
if "gemini_files" not in st.session_state:
    st.session_state.gemini_files = {}
if "messages" not in st.session_state:
    st.session_state.messages = []
if "generate_ppt" not in st.session_state:
    st.session_state.generate_ppt = False


# ── RAG Vector Store ───────────────────────────────────────────────────────────
@st.cache_resource
def load_vector_store():
    db_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "chroma_db")
    if os.path.exists(db_path):
        embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
        return Chroma(persist_directory=db_path, embedding_function=embeddings)
    return None

db = load_vector_store()


# ── AI Response ────────────────────────────────────────────────────────────────
def get_ai_response(user_prompt, class_level, subject, language, step_by_step, db):
    if not api_key:
        return "⚠️ Please configure `GEMINI_API_KEY` in your `.env` file."

    db_context = ""
    if db:
        try:
            docs = db.similarity_search(user_prompt, k=4)
            db_context = "\n".join([d.page_content for d in docs])
        except Exception:
            pass

    step_txt = "Always give a clear, numbered step-by-step breakdown." if step_by_step else ""
    system = f"""You are EduMitra, an expert AI tutor for Indian school and college students.
Help a Class {class_level} student with {subject}. Respond in: {language}.
Use the DATABASE CONTEXT and any UPLOADED FILES to give grounded answers.
Format answers with headers, bullet points, and helpful emojis. {step_txt}

--- DATABASE CONTEXT ---
{db_context}
--- END CONTEXT ---"""

    try:
        model = genai.GenerativeModel("gemini-2.5-flash")
        parts = [system]
        for gf in st.session_state.gemini_files.values():
            parts.append(gf)
        history = "\n".join([
            f"{m['role'].capitalize()}: {m['content']}"
            for m in st.session_state.messages[-6:]
        ])
        if history:
            parts.append(f"Recent conversation:\n{history}")
        parts.append(f"Student: {user_prompt}")
        return model.generate_content(parts).text
    except Exception as e:
        return f"⚠️ Error: {e}"


# ── PPT Helper ─────────────────────────────────────────────────────────────────
def create_pptx(slides_data):
    prs = Presentation()
    for s in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for b in s.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.level = 0
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# SIDEBAR
# ══════════════════════════════════════════════════════════════════════════════
with st.sidebar:
    # Brand + Theme Toggle
    toggle_icon = "☀️" if dark else "🌙"
    toggle_label = "Light Mode" if dark else "Dark Mode"
    brand_border = _t['sidebar_border']
    brand_color  = _t['primary']
    sub_color    = _t['text_secondary']
    st.markdown(f"""
    <div style="padding:24px 20px 20px; border-bottom:1px solid {brand_border};
                display:flex;justify-content:space-between;align-items:center;">
        <div>
            <div style="font-size:24px;font-weight:700;color:{brand_color};letter-spacing:-0.5px;">EduMitra</div>
            <div style="font-size:12px;color:{sub_color};margin-top:2px;">Intelligent Growth</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # Dark/Light toggle button
    if st.button(f"{toggle_icon}  {toggle_label}", key="theme_toggle"):
        st.session_state.dark_mode = not st.session_state.dark_mode
        st.rerun()

    # Nav links (visual only)
    nav_items = [
        #("📊", "Dashboard"),
       # ("📚", "My Courses"),
        ("🤖", "AI Tutor"),
       # ("📝", "Assignments"),
       # ("📈", "Analytics"),
    ]
    st.markdown('<div style="padding:16px 12px 8px;">', unsafe_allow_html=True)
    for icon, label in nav_items:
        active = label == "AI Tutor"
        bg     = _t['primary_light'] if active else "transparent"
        color  = _t['primary'] if active else _t['text_secondary']
        weight = "600" if active else "400"
        border = f"3px solid {_t['primary']}" if active else "3px solid transparent"
        st.markdown(f"""
        <div style="display:flex;align-items:center;gap:12px;padding:10px 14px;
                    border-radius:12px;background:{bg};border-right:{border};
                    margin-bottom:2px;cursor:pointer;">
            <span style="font-size:18px;">{icon}</span>
            <span style="font-size:14px;font-weight:{weight};color:{color};">{label}</span>
        </div>
        """, unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div style="padding:0 16px;">', unsafe_allow_html=True)
    st.markdown("---")

    # Tutor Preferences
    st.markdown(f'<p style="font-size:13px;font-weight:700;color:{_t["text_label"]};text-transform:uppercase;letter-spacing:0.06em;margin-bottom:12px;">Tutor Settings</p>', unsafe_allow_html=True)

    class_level = st.selectbox("Class Level", ["8", "9", "10", "11", "12", "College"])
    subject     = st.selectbox("Subject", ["Mathematics", "Physics", "Chemistry", "Biology", "Science", "General"])
    language    = st.selectbox("Language", ["English", "Tamil", "English + Tamil (Tamlish)"])
    step_by_step = st.checkbox("Step-by-step explanations", value=True)

    st.markdown("---")
    st.markdown(f'<p style="font-size:11px;color:{_t["text_secondary"]};line-height:1.5;">Powered by Gemini 2.5 Flash · 48,000+ NCERT chunks</p>', unsafe_allow_html=True)

    st.markdown("</div>", unsafe_allow_html=True)

    # Upgrade CTA
    st.markdown("""
    <div style="padding:16px;margin:12px 16px 20px;">
        <button style="width:100%;background:#006e2f;color:white;border:none;
                       border-radius:12px;padding:12px;font-size:14px;font-weight:600;
                       cursor:pointer;box-shadow:0 2px 8px rgba(0,110,47,0.2);">
            ✨ Upgrade to Pro
        </button>
    </div>
    """, unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# MAIN LAYOUT  (Chat | Right Panel)
# ══════════════════════════════════════════════════════════════════════════════
chat_col, right_col = st.columns([7, 3], gap="large")


# ── RIGHT PANEL ───────────────────────────────────────────────────────────────
with right_col:
    st.markdown("""
    <div style="padding:24px 0 0;">
        <p style="font-size:18px;font-weight:700;color:#0b1c30;margin:0 0 16px;">Knowledge Sources</p>
    </div>
    """, unsafe_allow_html=True)

    # File uploader
    uploaded_files = st.file_uploader(
        "Upload PDFs or images",
        type=["pdf", "png", "jpg", "jpeg"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )

    if st.button("⬆️  Upload & Process"):
        if not api_key:
            st.error("Add GEMINI_API_KEY to your .env file.")
        elif not uploaded_files:
            st.warning("Select at least one file first.")
        else:
            with st.spinner("Processing files…"):
                for uf in uploaded_files:
                    if uf.name not in st.session_state.gemini_files:
                        ext = "." + uf.name.split(".")[-1]
                        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                            tmp.write(uf.getvalue())
                            tmp_path = tmp.name
                        try:
                            gf = genai.upload_file(tmp_path, display_name=uf.name)
                            st.session_state.gemini_files[uf.name] = gf
                            st.success(f"✅ {uf.name}")
                        except Exception as e:
                            st.error(f"Failed: {e}")
                        finally:
                            os.remove(tmp_path)

    # Active sources list
    if st.session_state.gemini_files:
        st.markdown('<p style="font-size:12px;font-weight:700;color:#3d4a3d;text-transform:uppercase;letter-spacing:0.05em;margin:16px 0 8px;">Active Sources</p>', unsafe_allow_html=True)
        for name in st.session_state.gemini_files:
            ext = name.split(".")[-1].upper()
            icon_color = "#ba1a1a" if ext == "PDF" else "#494bd6"
            st.markdown(f"""
            <div style="display:flex;align-items:center;gap:10px;padding:10px 12px;
                        background:white;border:1px solid #dce9ff;border-radius:12px;margin-bottom:6px;">
                <span style="font-size:18px;color:{icon_color};">📄</span>
                <div style="min-width:0;">
                    <p style="font-size:13px;font-weight:500;color:#0b1c30;margin:0;
                               white-space:nowrap;overflow:hidden;text-overflow:ellipsis;">{name}</p>
                    <p style="font-size:11px;color:#6d7b6c;margin:0;">{ext} · uploaded</p>
                </div>
            </div>
            """, unsafe_allow_html=True)

        if st.button("🗑️  Clear All Sources"):
            for gf in st.session_state.gemini_files.values():
                try:
                    genai.delete_file(gf.name)
                except Exception:
                    pass
            st.session_state.gemini_files = {}
            st.rerun()

    st.markdown("---")

    # PPT Studio
    st.markdown('<p style="font-size:18px;font-weight:700;color:#0b1c30;margin:0 0 6px;">Presentation Studio</p>', unsafe_allow_html=True)
    st.markdown('<p style="font-size:13px;color:#6d7b6c;margin:0 0 14px;">Generate a PowerPoint from your documents & chat.</p>', unsafe_allow_html=True)

    if st.button("✨  Create Summary PPT"):
        st.session_state.generate_ppt = True

    if st.session_state.generate_ppt:
        with st.spinner("Crafting slides…"):
            try:
                model = genai.GenerativeModel("gemini-2.5-flash")
                prompt = """Summarize key concepts into a presentation.
Output ONLY valid JSON (no markdown fences):
[{"title": "...", "bullets": ["...", "..."]}]
Create 5–10 educational slides."""
                parts = [prompt]
                for gf in st.session_state.gemini_files.values():
                    parts.append(gf)
                hist = "\n".join([f"{m['role'].capitalize()}: {m['content']}" for m in st.session_state.messages])
                if hist:
                    parts.append(f"Conversation:\n{hist}")
                raw = model.generate_content(parts).text.strip()
                for marker in ["```json", "```"]:
                    raw = raw.replace(marker, "")
                slides = json.loads(raw)
                ppt_buf = create_pptx(slides)
                st.success("Presentation ready!")
                st.download_button(
                    "📥  Download .pptx",
                    data=ppt_buf,
                    file_name="EduMitra_Summary.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            except Exception as e:
                st.error(f"Failed: {e}")
        st.session_state.generate_ppt = False


# ── CHAT PANEL ────────────────────────────────────────────────────────────────
with chat_col:
    # Top bar
    ncert_loaded = bool(db)
    sources_loaded = bool(st.session_state.gemini_files)
    badge_color = "#006e2f" if (ncert_loaded or sources_loaded) else "#6d7b6c"
    badge_text  = "NCERT Context Loaded" if ncert_loaded else "No RAG Context"

    st.markdown(f"""
    <div style="display:flex;justify-content:space-between;align-items:center;
                padding:20px 4px 16px;border-bottom:1px solid #dce9ff;margin-bottom:20px;">
        <h2 style="margin:0;font-size:22px;font-weight:700;color:#006e2f;">AI Tutor</h2>
        <span style="display:inline-flex;align-items:center;gap:6px;background:#f0fdf4;
                     border:1px solid #bbf7d0;border-radius:20px;padding:5px 12px;
                     font-size:12px;font-weight:600;color:{badge_color};">
            ● {badge_text}
        </span>
    </div>
    """, unsafe_allow_html=True)

    # Welcome message
    if not st.session_state.messages:
        st.markdown("""
        <div style="background:rgba(240,253,244,0.9);border:1px solid rgba(74,225,118,0.2);
                    border-radius:20px 20px 20px 4px;padding:20px 24px;margin-bottom:24px;
                    box-shadow:0 1px 4px rgba(0,0,0,0.04);">
            <p style="font-size:16px;font-weight:600;color:#006e2f;margin:0 0 8px;">
                👋 Hello! I'm EduMitra.
            </p>
            <p style="font-size:15px;color:#0b1c30;margin:0 0 16px;line-height:1.6;">
                I'm your personalised AI tutor for CBSE & State Board subjects. 
                Ask me anything — I'll explain concepts step by step, in your language.
            </p>
            <div style="display:flex;flex-wrap:wrap;gap:8px;">
                <span style="background:white;border:1px solid #dce9ff;border-radius:20px;
                             padding:6px 14px;font-size:13px;color:#006e2f;font-weight:500;">
                    ✨ Explain step-by-step
                </span>
                <span style="background:white;border:1px solid #dce9ff;border-radius:20px;
                             padding:6px 14px;font-size:13px;color:#006e2f;font-weight:500;">
                    🧪 Generate practice questions
                </span>
                <span style="background:white;border:1px solid #dce9ff;border-radius:20px;
                             padding:6px 14px;font-size:13px;color:#006e2f;font-weight:500;">
                    🌐 Answer in Tamil
                </span>
            </div>
        </div>
        """, unsafe_allow_html=True)

    # Chat history
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    # Chat input
    if user_input := st.chat_input("Ask EduMitra anything…"):
        st.session_state.messages.append({"role": "user", "content": user_input})
        with st.chat_message("user"):
            st.markdown(user_input)

        with st.chat_message("assistant"):
            with st.spinner("EduMitra is thinking…"):
                reply = get_ai_response(user_input, class_level, subject, language, step_by_step, db)
            st.markdown(reply)

        st.session_state.messages.append({"role": "assistant", "content": reply})
